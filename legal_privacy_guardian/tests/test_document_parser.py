import unittest
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from pptx import Presentation

from document_parser import DocumentExtractionError, extract_document_text
from privacy import mask_sensitive_data


EMAIL = "test@example.com"


def make_docx() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "word/document.xml",
            f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{EMAIL}</w:t></w:r></w:p><w:sectPr/></w:body></w:document>",
        )
    return output.getvalue()


def make_pdf(text: str = EMAIL) -> bytes:
    stream = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF".encode()
    )
    return bytes(pdf)


def make_xlsx() -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = EMAIL
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 5_000_000, 500_000).text = EMAIL
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class DocumentParserTests(unittest.TestCase):
    def test_extracts_sensitive_text_from_supported_formats(self):
        samples = {
            ".txt": EMAIL.encode(),
            ".md": EMAIL.encode(),
            ".html": f"<p>{EMAIL}</p>".encode(),
            ".csv": f"email\n{EMAIL}".encode(),
            ".json": f'{{"email":"{EMAIL}"}}'.encode(),
            ".xml": f"<email>{EMAIL}</email>".encode(),
            ".rtf": f"{{\\rtf1\\ansi {EMAIL}}}".encode(),
            ".pdf": make_pdf(),
            ".docx": make_docx(),
            ".xlsx": make_xlsx(),
            ".pptx": make_pptx(),
        }

        for extension, content in samples.items():
            with self.subTest(extension=extension):
                extracted = extract_document_text(f"sample{extension}", content)
                findings, redacted = mask_sensitive_data(extracted)
                self.assertEqual(findings["adresses e-mail"], 1)
                self.assertNotIn(EMAIL, redacted)

    def test_rejects_unknown_extensions_and_scanned_image_only_pdf(self):
        with self.assertRaises(DocumentExtractionError):
            extract_document_text("photo.png", b"not an image")

        image_only_pdf = make_pdf("")
        with self.assertRaises(DocumentExtractionError):
            extract_document_text("scan.pdf", image_only_pdf)


if __name__ == "__main__":
    unittest.main()
